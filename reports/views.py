"""
Rapports & Statistiques — toutes les données proviennent de MySQL
(section 15 du cahier des charges), aucune valeur codée en dur.
"""
from datetime import timedelta
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.db.models import Sum
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone

from core.charts import bar_heights_percent, line_chart_svg
from core.permissions import VIEW_REPORTS, role_required
from products.models import Category, Product
from sales.models import Sale, SaleItem


def _gather_report_data():
    today = timezone.localdate()
    start = today - timedelta(days=29)

    products = Product.objects.all()
    low_stock_qs = [p for p in products if p.is_low_stock or p.is_out_of_stock]
    inventory_value = sum((p.stock_value for p in products), Decimal('0'))

    # Chiffre d'affaires quotidien sur 30 jours (données réelles de Sale)
    daily_totals = []
    day_labels = []
    for i in range(30):
        day = start + timedelta(days=i)
        total = Sale.objects.filter(date__date=day, status=Sale.Status.COMPLETED).aggregate(t=Sum('total'))['t'] or 0
        daily_totals.append(float(total))
        day_labels.append(day.strftime('%d/%m'))

    chart = line_chart_svg(daily_totals or [0, 0])

    # Ventes par catégorie (montant total des lignes de vente, toutes dates)
    category_sales = (
        SaleItem.objects
        .values('product__category__name')
        .annotate(total=Sum('subtotal'))
        .order_by('-total')[:6]
    )
    cat_labels = [c['product__category__name'] or 'Sans catégorie' for c in category_sales]
    cat_values = [float(c['total'] or 0) for c in category_sales]
    cat_heights = bar_heights_percent(cat_values)
    cat_colors = ['#0058be', '#38bdf8', '#4edea3', '#f59e0b', '#a78bfa', '#f87171']

    return {
        'low_stock_count': len(low_stock_qs),
        'low_stock_products': low_stock_qs[:10],
        'inventory_value': inventory_value,
        'chart': chart,
        'day_labels': day_labels,
        'category_bars': list(zip(cat_labels, cat_values, cat_heights, cat_colors)),
        'total_products': products.count(),
        'total_categories': Category.objects.count(),
        'total_revenue_30d': sum(daily_totals),
    }


@role_required(*VIEW_REPORTS)
def index(request):
    return render(request, 'reports/index.html', _gather_report_data())


@role_required(*VIEW_REPORTS)
def export_stock_csv(request):
    import csv
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="rapport_stock.csv"'
    writer = csv.writer(response)
    writer.writerow(['Produit', 'Référence', 'Catégorie', 'Stock', 'Seuil min.', 'Valeur du stock'])
    for p in Product.objects.select_related('category'):
        writer.writerow([p.name, p.reference, p.category.name if p.category else '', p.stock_quantity, p.minimum_stock, p.stock_value])
    return response


@role_required(*VIEW_REPORTS)
def export_stock_pdf(request):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    import io

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = [Paragraph("Rapport de stock — Kamsar Street", styles['Title']), Spacer(1, 12)]

    data = [['Produit', 'Référence', 'Catégorie', 'Stock', 'Seuil min.', 'Valeur (GNF)']]
    for p in Product.objects.select_related('category'):
        data.append([p.name, p.reference, p.category.name if p.category else '—', str(p.stock_quantity), str(p.minimum_stock), f"{p.stock_value:,.0f}"])

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#131b2e')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F1F5F9')]),
    ]))
    elements.append(table)
    doc.build(elements)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="rapport_stock.pdf"'
    return response


@role_required(*VIEW_REPORTS)
def export_stock_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    import io

    wb = Workbook()
    ws = wb.active
    ws.title = 'Stock'
    headers = ['Produit', 'Référence', 'Catégorie', 'Stock', 'Seuil min.', 'Valeur du stock (GNF)']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='131B2E')

    for p in Product.objects.select_related('category'):
        ws.append([p.name, p.reference, p.category.name if p.category else '', p.stock_quantity, p.minimum_stock, float(p.stock_value)])

    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="rapport_stock.xlsx"'
    return response
